from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    Border,
    Side,
    PatternFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break


class ExportadorExcelService:

    def __init__(
        self,
        pasta="exports",
        linhas_corpo_impressao=44
    ):
        self._pasta = Path(pasta)

        self._pasta.mkdir(
            parents=True,
            exist_ok=True
        )

        self._linhas_corpo_impressao = (
            linhas_corpo_impressao
        )

    def exportar(
        self,
        escala,
        pessoas
    ):

        nome_arquivo = self._nome_arquivo(
            escala.nome
        )

        caminho = self._pasta / nome_arquivo

        workbook = Workbook()

        # ==================================================
        # ABA ESCALA
        # ==================================================

        planilha = workbook.active
        planilha.title = "ESCALA"

        self._criar_conteudo(
            planilha,
            escala
        )

        self._formatar(
            planilha
        )

        self._adicionar_contatos(
            planilha,
            escala,
            pessoas
        )

        self._ajustar_area_impressao(
            planilha
        )

        # ==================================================
        # ABA DISTRIBUIÇÃO
        # ==================================================

        planilha_distribuicao = (
            workbook.create_sheet(
                "DISTRIBUIÇÃO"
            )
        )

        self._criar_distribuicao(
            planilha_distribuicao,
            escala
        )

        # ==================================================
        # ABA IMPRESSÃO
        # ==================================================

        planilha_impressao = (
            workbook.create_sheet(
                "IMPRESSÃO"
            )
        )

        self._criar_impressao(
            planilha_impressao,
            escala,
            pessoas
        )

        workbook.save(
            caminho
        )

        return caminho

    # ==================================================
    # CONTEÚDO DA ESCALA
    # ==================================================

    def _criar_conteudo(
        self,
        planilha,
        escala
    ):

        planilha.merge_cells(
            "A1:D1"
        )

        planilha["A1"] = (
            "Congregação Cristã no Brasil"
        )

        planilha.merge_cells(
            "A2:D2"
        )

        planilha["A2"] = escala.nome

        planilha.merge_cells(
            "A3:D3"
        )

        planilha["A3"] = (
            f"PERÍODO: "
            f"{self._obter_periodo(escala)}"
        )

        planilha.append([])

        planilha.append([
            "DATA",
            "EVENTO",
            "ATENDENTE",
            "FUNÇÃO"
        ])

        for ocorrencia in escala.ocorrencias:

            for atribuicao in (
                ocorrencia.atribuicoes
            ):

                planilha.append([
                    ocorrencia.data,
                    ocorrencia.evento.nome,
                    atribuicao["pessoa"],
                    atribuicao["funcao"]
                ])

    # ==================================================
    # FORMATAÇÃO DA ESCALA
    # ==================================================

    def _formatar(
        self,
        planilha
    ):

        planilha.sheet_view.showGridLines = False

        preenchimento_titulo = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        preenchimento_periodo = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="5B9BD5"
        )

        preenchimento_azul_claro = PatternFill(
            fill_type="solid",
            fgColor="DDEBF7"
        )

        preenchimento_branco = PatternFill(
            fill_type="solid",
            fgColor="FFFFFF"
        )

        lado_fino = Side(
            style="thin",
            color="B7B7B7"
        )

        lado_medio = Side(
            style="medium",
            color="404040"
        )

        # --------------------------------------------------
        # TÍTULO PRINCIPAL
        # --------------------------------------------------

        planilha["A1"].font = Font(
            bold=True,
            size=16,
            color="FFFFFF"
        )

        planilha["A1"].fill = (
            preenchimento_titulo
        )

        planilha["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=1,
            linha_fim=1,
            coluna_inicio=1,
            coluna_fim=4,
            lado=lado_medio
        )

        planilha.row_dimensions[1].height = 26

        # --------------------------------------------------
        # NOME DA ESCALA
        # --------------------------------------------------

        planilha["A2"].font = Font(
            bold=True,
            size=14,
            color="FFFFFF"
        )

        planilha["A2"].fill = (
            preenchimento_titulo
        )

        planilha["A2"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=2,
            linha_fim=2,
            coluna_inicio=1,
            coluna_fim=4,
            lado=lado_medio
        )

        planilha.row_dimensions[2].height = 24

        # --------------------------------------------------
        # PERÍODO
        # --------------------------------------------------

        planilha["A3"].font = Font(
            bold=True,
            size=10
        )

        planilha["A3"].fill = (
            preenchimento_periodo
        )

        planilha["A3"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=3,
            linha_fim=3,
            coluna_inicio=1,
            coluna_fim=4,
            lado=lado_medio
        )

        planilha.row_dimensions[3].height = 21

        # --------------------------------------------------
        # CABEÇALHO
        # --------------------------------------------------

        for celula in planilha[5]:

            celula.font = Font(
                bold=True,
                color="FFFFFF"
            )

            celula.fill = (
                preenchimento_cabecalho
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            celula.border = Border(
                top=lado_medio,
                bottom=lado_medio,
                left=lado_fino,
                right=lado_fino
            )

        planilha.row_dimensions[5].height = 22

        # --------------------------------------------------
        # DADOS
        # --------------------------------------------------

        for linha in planilha.iter_rows(
            min_row=6,
            max_row=planilha.max_row,
            min_col=1,
            max_col=4
        ):

            for celula in linha:

                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                celula.border = Border(
                    top=lado_fino,
                    bottom=lado_fino,
                    left=lado_fino,
                    right=lado_fino
                )

        # --------------------------------------------------
        # ATENDENTE
        # --------------------------------------------------

        for linha in range(
            6,
            planilha.max_row + 1
        ):

            planilha.cell(
                row=linha,
                column=3
            ).alignment = Alignment(
                horizontal="left",
                vertical="center"
            )

        # --------------------------------------------------
        # DATA
        # --------------------------------------------------

        for linha in range(
            6,
            planilha.max_row + 1
        ):

            celula = planilha.cell(
                row=linha,
                column=1
            )

            celula.number_format = (
                "dd/mm/yyyy"
            )

        # --------------------------------------------------
        # GRUPOS
        # --------------------------------------------------

        self._formatar_grupos(
            planilha,
            lado_medio,
            preenchimento_azul_claro,
            preenchimento_branco
        )

        self._mesclar_grupos(
            planilha
        )

        # --------------------------------------------------
        # LARGURA
        # --------------------------------------------------

        self._ajustar_largura_colunas(
            planilha
        )

        # --------------------------------------------------
        # ALTURA
        # --------------------------------------------------

        for linha in range(
            6,
            planilha.max_row + 1
        ):

            planilha.row_dimensions[
                linha
            ].height = 18

        # --------------------------------------------------
        # CONGELAMENTO
        # --------------------------------------------------

        planilha.freeze_panes = "A6"

        # --------------------------------------------------
        # IMPRESSÃO
        # --------------------------------------------------

        planilha.page_setup.orientation = (
            "landscape"
        )

        planilha.page_setup.paperSize = (
            planilha.PAPERSIZE_A4
        )

        planilha.page_setup.fitToWidth = 1
        planilha.page_setup.fitToHeight = 0

        planilha.sheet_properties.pageSetUpPr.fitToPage = (
            True
        )

        planilha.print_title_rows = "1:5"

        planilha.page_margins.left = 0.25
        planilha.page_margins.right = 0.25
        planilha.page_margins.top = 0.5
        planilha.page_margins.bottom = 0.5

    # ==================================================
    # CONTATOS DA ABA ESCALA
    # ==================================================

    def _adicionar_contatos(
        self,
        planilha,
        escala,
        pessoas
    ):

        contatos = self._obter_contatos_escalados(
            escala,
            pessoas
        )

        if not contatos:
            return

        linha_titulo = (
            planilha.max_row + 2
        )

        preenchimento_titulo = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="5B9BD5"
        )

        preenchimento_azul_claro = PatternFill(
            fill_type="solid",
            fgColor="DDEBF7"
        )

        preenchimento_branco = PatternFill(
            fill_type="solid",
            fgColor="FFFFFF"
        )

        lado_fino = Side(
            style="thin",
            color="B7B7B7"
        )

        lado_medio = Side(
            style="medium",
            color="404040"
        )

        # --------------------------------------------------
        # TÍTULO DOS CONTATOS
        # --------------------------------------------------

        planilha.merge_cells(
            start_row=linha_titulo,
            start_column=1,
            end_row=linha_titulo,
            end_column=4
        )

        titulo = planilha.cell(
            row=linha_titulo,
            column=1
        )

        titulo.value = (
            "CONTATOS DOS ATENDENTES"
        )

        titulo.font = Font(
            bold=True,
            color="FFFFFF",
            size=11
        )

        titulo.fill = (
            preenchimento_titulo
        )

        titulo.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha_titulo,
            linha_fim=linha_titulo,
            coluna_inicio=1,
            coluna_fim=4,
            lado=lado_medio
        )

        planilha.row_dimensions[
            linha_titulo
        ].height = 20

        # --------------------------------------------------
        # CABEÇALHO
        # --------------------------------------------------

        linha_cabecalho = (
            linha_titulo + 1
        )

        planilha.merge_cells(
            start_row=linha_cabecalho,
            start_column=1,
            end_row=linha_cabecalho,
            end_column=2
        )

        planilha.merge_cells(
            start_row=linha_cabecalho,
            start_column=3,
            end_row=linha_cabecalho,
            end_column=4
        )

        celula_nome = planilha.cell(
            row=linha_cabecalho,
            column=1
        )

        celula_nome.value = "NOME"

        celula_telefone = planilha.cell(
            row=linha_cabecalho,
            column=3
        )

        celula_telefone.value = (
            "TELEFONE"
        )

        for celula in (
            celula_nome,
            celula_telefone
        ):

            celula.font = Font(
                bold=True,
                color="FFFFFF"
            )

            celula.fill = (
                preenchimento_cabecalho
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha_cabecalho,
            linha_fim=linha_cabecalho,
            coluna_inicio=1,
            coluna_fim=2,
            lado=lado_medio
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha_cabecalho,
            linha_fim=linha_cabecalho,
            coluna_inicio=3,
            coluna_fim=4,
            lado=lado_medio
        )

        planilha.row_dimensions[
            linha_cabecalho
        ].height = 18

        # --------------------------------------------------
        # LISTA
        # --------------------------------------------------

        linha_atual = (
            linha_cabecalho + 1
        )

        for indice, pessoa in enumerate(
            contatos
        ):

            planilha.merge_cells(
                start_row=linha_atual,
                start_column=1,
                end_row=linha_atual,
                end_column=2
            )

            planilha.merge_cells(
                start_row=linha_atual,
                start_column=3,
                end_row=linha_atual,
                end_column=4
            )

            nome = planilha.cell(
                row=linha_atual,
                column=1
            )

            telefone = planilha.cell(
                row=linha_atual,
                column=3
            )

            nome.value = pessoa.nome
            telefone.value = pessoa.telefone

            if indice % 2 == 0:

                preenchimento = (
                    preenchimento_azul_claro
                )

            else:

                preenchimento = (
                    preenchimento_branco
                )

            nome.fill = preenchimento
            telefone.fill = preenchimento

            nome.alignment = Alignment(
                horizontal="left",
                vertical="center"
            )

            telefone.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            self._aplicar_borda_intervalo(
                planilha,
                linha_inicio=linha_atual,
                linha_fim=linha_atual,
                coluna_inicio=1,
                coluna_fim=2,
                lado=lado_fino
            )

            self._aplicar_borda_intervalo(
                planilha,
                linha_inicio=linha_atual,
                linha_fim=linha_atual,
                coluna_inicio=3,
                coluna_fim=4,
                lado=lado_fino
            )

            planilha.row_dimensions[
                linha_atual
            ].height = 18

            linha_atual += 1

        ultima_linha = (
            linha_atual - 1
        )

        self._aplicar_borda_externa(
            planilha,
            linha_inicio=linha_cabecalho,
            linha_fim=ultima_linha,
            coluna_inicio=1,
            coluna_fim=4,
            lado=lado_medio
        )

    # ==================================================
    # DISTRIBUIÇÃO
    # ==================================================

    def _criar_distribuicao(
        self,
        planilha,
        escala
    ):

        planilha.sheet_view.showGridLines = False

        funcoes = (
            self._obter_funcoes_distribuicao(
                escala
            )
        )

        pessoas = (
            self._obter_pessoas_distribuicao(
                escala
            )
        )

        ultima_coluna = (
            len(funcoes) + 2
        )

        planilha.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=ultima_coluna
        )

        planilha["A1"] = (
            "Congregação Cristã no Brasil"
        )

        planilha.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=ultima_coluna
        )

        planilha["A2"] = (
            f"DISTRIBUIÇÃO - {escala.nome}"
        )

        planilha.merge_cells(
            start_row=3,
            start_column=1,
            end_row=3,
            end_column=ultima_coluna
        )

        planilha["A3"] = (
            f"PERÍODO: "
            f"{self._obter_periodo(escala)}"
        )

        self._formatar_titulo_distribuicao(
            planilha,
            ultima_coluna
        )

        linha_atual = 5

        # --------------------------------------------------
        # TOTAL DO PERÍODO
        # --------------------------------------------------

        contagem_total = (
            self._contar_distribuicao(
                escala.ocorrencias
            )
        )

        linha_atual = (
            self._adicionar_tabela_distribuicao(
                planilha=planilha,
                linha_inicio=linha_atual,
                titulo="TOTAL DO PERÍODO",
                pessoas=pessoas,
                funcoes=funcoes,
                contagem=contagem_total
            )
        )

        # --------------------------------------------------
        # DISTRIBUIÇÃO MENSAL
        # --------------------------------------------------

        ocorrencias_por_mes = (
            self._agrupar_ocorrencias_por_mes(
                escala
            )
        )

        for chave_mes in sorted(
            ocorrencias_por_mes
        ):

            ocorrencias_mes = (
                ocorrencias_por_mes[
                    chave_mes
                ]
            )

            contagem_mes = (
                self._contar_distribuicao(
                    ocorrencias_mes
                )
            )

            pessoas_mes = (
                self._obter_pessoas_das_ocorrencias(
                    ocorrencias_mes
                )
            )

            nome_mes = (
                self._nome_mes(
                    chave_mes[1]
                )
            )

            titulo_mes = (
                f"{nome_mes}/{chave_mes[0]}"
            )

            linha_atual = (
                self._adicionar_tabela_distribuicao(
                    planilha=planilha,
                    linha_inicio=linha_atual,
                    titulo=titulo_mes,
                    pessoas=pessoas_mes,
                    funcoes=funcoes,
                    contagem=contagem_mes
                )
            )

        self._ajustar_largura_distribuicao(
            planilha,
            ultima_coluna
        )

        planilha.freeze_panes = "A5"

        planilha.page_setup.orientation = (
            "landscape"
        )

        planilha.page_setup.paperSize = (
            planilha.PAPERSIZE_A4
        )

        planilha.page_setup.fitToWidth = 1
        planilha.page_setup.fitToHeight = 0

        planilha.sheet_properties.pageSetUpPr.fitToPage = (
            True
        )

        planilha.page_margins.left = 0.25
        planilha.page_margins.right = 0.25
        planilha.page_margins.top = 0.5
        planilha.page_margins.bottom = 0.5

        planilha.print_area = (
            f"A1:"
            f"{get_column_letter(ultima_coluna)}"
            f"{planilha.max_row}"
        )

    def _formatar_titulo_distribuicao(
        self,
        planilha,
        ultima_coluna
    ):

        preenchimento_titulo = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        preenchimento_periodo = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        lado_medio = Side(
            style="medium",
            color="404040"
        )

        for linha in (
            1,
            2
        ):

            celula = planilha.cell(
                row=linha,
                column=1
            )

            celula.fill = (
                preenchimento_titulo
            )

            celula.font = Font(
                bold=True,
                color="FFFFFF",
                size=16 if linha == 1 else 14
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            self._aplicar_borda_intervalo(
                planilha,
                linha_inicio=linha,
                linha_fim=linha,
                coluna_inicio=1,
                coluna_fim=ultima_coluna,
                lado=lado_medio
            )

        planilha["A3"].fill = (
            preenchimento_periodo
        )

        planilha["A3"].font = Font(
            bold=True,
            size=10
        )

        planilha["A3"].alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=3,
            linha_fim=3,
            coluna_inicio=1,
            coluna_fim=ultima_coluna,
            lado=lado_medio
        )

        planilha.row_dimensions[1].height = 26
        planilha.row_dimensions[2].height = 24
        planilha.row_dimensions[3].height = 21

    def _adicionar_tabela_distribuicao(
        self,
        planilha,
        linha_inicio,
        titulo,
        pessoas,
        funcoes,
        contagem
    ):

        ultima_coluna = (
            len(funcoes) + 2
        )

        preenchimento_titulo = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="5B9BD5"
        )

        preenchimento_azul_claro = PatternFill(
            fill_type="solid",
            fgColor="DDEBF7"
        )

        preenchimento_branco = PatternFill(
            fill_type="solid",
            fgColor="FFFFFF"
        )

        lado_fino = Side(
            style="thin",
            color="B7B7B7"
        )

        lado_medio = Side(
            style="medium",
            color="404040"
        )

        # --------------------------------------------------
        # TÍTULO
        # --------------------------------------------------

        planilha.merge_cells(
            start_row=linha_inicio,
            start_column=1,
            end_row=linha_inicio,
            end_column=ultima_coluna
        )

        celula_titulo = planilha.cell(
            row=linha_inicio,
            column=1
        )

        celula_titulo.value = titulo

        celula_titulo.fill = (
            preenchimento_titulo
        )

        celula_titulo.font = Font(
            bold=True,
            color="FFFFFF",
            size=11
        )

        celula_titulo.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha_inicio,
            linha_fim=linha_inicio,
            coluna_inicio=1,
            coluna_fim=ultima_coluna,
            lado=lado_medio
        )

        planilha.row_dimensions[
            linha_inicio
        ].height = 20

        # --------------------------------------------------
        # CABEÇALHO
        # --------------------------------------------------

        linha_cabecalho = (
            linha_inicio + 1
        )

        planilha.cell(
            row=linha_cabecalho,
            column=1
        ).value = "PESSOA"

        for indice, funcao in enumerate(
            funcoes,
            start=2
        ):

            planilha.cell(
                row=linha_cabecalho,
                column=indice
            ).value = funcao

        planilha.cell(
            row=linha_cabecalho,
            column=ultima_coluna
        ).value = "TOTAL"

        for coluna in range(
            1,
            ultima_coluna + 1
        ):

            celula = planilha.cell(
                row=linha_cabecalho,
                column=coluna
            )

            celula.fill = (
                preenchimento_cabecalho
            )

            celula.font = Font(
                bold=True,
                color="FFFFFF"
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            celula.border = Border(
                top=lado_medio,
                bottom=lado_medio,
                left=lado_fino,
                right=lado_fino
            )

        planilha.row_dimensions[
            linha_cabecalho
        ].height = 18

        # --------------------------------------------------
        # DADOS
        # --------------------------------------------------

        linha_atual = (
            linha_cabecalho + 1
        )

        for indice_pessoa, pessoa in enumerate(
            pessoas
        ):

            if indice_pessoa % 2 == 0:

                preenchimento = (
                    preenchimento_azul_claro
                )

            else:

                preenchimento = (
                    preenchimento_branco
                )

            planilha.cell(
                row=linha_atual,
                column=1
            ).value = pessoa

            total_pessoa = 0

            for indice_funcao, funcao in enumerate(
                funcoes,
                start=2
            ):

                quantidade = (
                    contagem
                    .get(pessoa, {})
                    .get(funcao, 0)
                )

                total_pessoa += quantidade

                planilha.cell(
                    row=linha_atual,
                    column=indice_funcao
                ).value = quantidade

            planilha.cell(
                row=linha_atual,
                column=ultima_coluna
            ).value = total_pessoa

            for coluna in range(
                1,
                ultima_coluna + 1
            ):

                celula = planilha.cell(
                    row=linha_atual,
                    column=coluna
                )

                celula.fill = (
                    preenchimento
                )

                celula.border = Border(
                    top=lado_fino,
                    bottom=lado_fino,
                    left=lado_fino,
                    right=lado_fino
                )

                if coluna == 1:

                    celula.alignment = Alignment(
                        horizontal="left",
                        vertical="center"
                    )

                else:

                    celula.alignment = Alignment(
                        horizontal="center",
                        vertical="center"
                    )

            planilha.cell(
                row=linha_atual,
                column=ultima_coluna
            ).font = Font(
                bold=True
            )

            planilha.row_dimensions[
                linha_atual
            ].height = 18

            linha_atual += 1

        # --------------------------------------------------
        # TOTAL
        # --------------------------------------------------

        linha_total = linha_atual

        planilha.cell(
            row=linha_total,
            column=1
        ).value = "TOTAL"

        total_geral = 0

        for indice_funcao, funcao in enumerate(
            funcoes,
            start=2
        ):

            total_funcao = 0

            for pessoa in pessoas:

                total_funcao += (
                    contagem
                    .get(pessoa, {})
                    .get(funcao, 0)
                )

            total_geral += total_funcao

            planilha.cell(
                row=linha_total,
                column=indice_funcao
            ).value = total_funcao

        planilha.cell(
            row=linha_total,
            column=ultima_coluna
        ).value = total_geral

        for coluna in range(
            1,
            ultima_coluna + 1
        ):

            celula = planilha.cell(
                row=linha_total,
                column=coluna
            )

            celula.font = Font(
                bold=True
            )

            celula.fill = (
                preenchimento_cabecalho
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            celula.border = Border(
                top=lado_medio,
                bottom=lado_medio,
                left=lado_fino,
                right=lado_fino
            )

        planilha.row_dimensions[
            linha_total
        ].height = 18

        self._aplicar_borda_externa(
            planilha,
            linha_inicio=linha_cabecalho,
            linha_fim=linha_total,
            coluna_inicio=1,
            coluna_fim=ultima_coluna,
            lado=lado_medio
        )

        return linha_total + 3

    # ==================================================
    # ABA IMPRESSÃO
    # ==================================================

    def _criar_impressao(
        self,
        planilha,
        escala,
        pessoas
    ):

        planilha.sheet_view.showGridLines = False

        paginas_logicas = (
            self._paginar_impressao(
                escala,
                pessoas
            )
        )

        faces = (
            self._montar_faces_impressao(
                paginas_logicas
            )
        )

        altura_face = (
            self._altura_face_impressao()
        )

        for indice_face, face in enumerate(
            faces
        ):

            linha_inicio = (
                indice_face * altura_face
                + 1
            )

            pagina_esquerda = face[
                "esquerda"
            ]

            pagina_direita = face[
                "direita"
            ]

            if pagina_esquerda is not None:

                self._renderizar_pagina_logica(
                    planilha=planilha,
                    escala=escala,
                    pagina=pagina_esquerda,
                    linha_inicio=linha_inicio,
                    coluna_inicio=1
                )

            if pagina_direita is not None:

                self._renderizar_pagina_logica(
                    planilha=planilha,
                    escala=escala,
                    pagina=pagina_direita,
                    linha_inicio=linha_inicio,
                    coluna_inicio=6
                )

            linha_fim_face = (
                linha_inicio
                + altura_face
                - 1
            )

            for linha in range(
                linha_inicio,
                linha_fim_face + 1
            ):

                if (
                    planilha.row_dimensions[
                        linha
                    ].height
                    is None
                ):
                    planilha.row_dimensions[
                        linha
                    ].height = 17

            if indice_face < len(faces) - 1:

                planilha.row_breaks.append(
                    Break(
                        id=linha_fim_face
                    )
                )

        self._ajustar_colunas_impressao(
            planilha
        )

        ultima_linha = (
            len(faces) * altura_face
        )

        # --------------------------------------------------
        # CONFIGURAÇÃO DE IMPRESSÃO
        # --------------------------------------------------

        planilha.page_setup.orientation = (
            "portrait"
        )

        planilha.page_setup.paperSize = (
            planilha.PAPERSIZE_A4
        )

        planilha.page_setup.fitToWidth = 1
        planilha.page_setup.fitToHeight = 0

        planilha.page_setup.pageOrder = (
            "downThenOver"
        )

        planilha.sheet_properties.pageSetUpPr.fitToPage = (
            True
        )

        planilha.page_margins.left = 0.15
        planilha.page_margins.right = 0.15
        planilha.page_margins.top = 0.20
        planilha.page_margins.bottom = 0.20
        planilha.page_margins.header = 0
        planilha.page_margins.footer = 0

        planilha.print_area = (
            f"A1:I{ultima_linha}"
        )

    # ==================================================
    # PAGINAÇÃO DA IMPRESSÃO
    # ==================================================

    def _paginar_impressao(
        self,
        escala,
        pessoas
    ):

        capacidade = (
            self._linhas_corpo_impressao
        )

        paginas = []

        pagina_atual = {
            "eventos": [],
            "contatos": []
        }

        linhas_usadas = 0

        # --------------------------------------------------
        # EVENTOS
        # --------------------------------------------------

        for ocorrencia in escala.ocorrencias:

            linhas_evento = max(
                1,
                len(
                    ocorrencia.atribuicoes
                )
            )

            # Primeira ocorrência da página precisa também
            # da linha do cabeçalho da escala.
            if not pagina_atual["eventos"]:

                linhas_necessarias = (
                    1 + linhas_evento
                )

            else:

                linhas_necessarias = (
                    linhas_evento
                )

            if linhas_necessarias > capacidade:

                raise ValueError(
                    "Um evento possui mais linhas do que "
                    "a capacidade de uma página de impressão."
                )

            if (
                linhas_usadas
                + linhas_necessarias
                > capacidade
            ):

                paginas.append(
                    pagina_atual
                )

                pagina_atual = {
                    "eventos": [],
                    "contatos": []
                }

                linhas_usadas = 0

                linhas_necessarias = (
                    1 + linhas_evento
                )

            pagina_atual[
                "eventos"
            ].append(
                ocorrencia
            )

            linhas_usadas += (
                linhas_necessarias
            )

        # --------------------------------------------------
        # CONTATOS
        # --------------------------------------------------

        contatos = (
            self._obter_contatos_escalados(
                escala,
                pessoas
            )
        )

        if contatos:

            # Título + cabeçalho + pessoas
            linhas_contatos = (
                2 + len(contatos)
            )

            if linhas_contatos > capacidade:

                raise ValueError(
                    "A relação de contatos é maior do que "
                    "a capacidade de uma página de impressão."
                )

            if pagina_atual["eventos"]:

                # Uma linha em branco antes dos contatos.
                linhas_necessarias = (
                    1 + linhas_contatos
                )

            else:

                linhas_necessarias = (
                    linhas_contatos
                )

            if (
                linhas_usadas
                + linhas_necessarias
                <= capacidade
            ):

                pagina_atual[
                    "contatos"
                ] = contatos

                linhas_usadas += (
                    linhas_necessarias
                )

            else:

                if (
                    pagina_atual["eventos"]
                    or pagina_atual["contatos"]
                ):

                    paginas.append(
                        pagina_atual
                    )

                pagina_atual = {
                    "eventos": [],
                    "contatos": contatos
                }

                linhas_usadas = (
                    linhas_contatos
                )

        # --------------------------------------------------
        # ÚLTIMA PÁGINA
        # --------------------------------------------------

        if (
            pagina_atual["eventos"]
            or pagina_atual["contatos"]
        ):

            paginas.append(
                pagina_atual
            )

        return paginas

    # ==================================================
    # IMPOSIÇÃO FRENTE / VERSO
    # ==================================================

    def _montar_faces_impressao(
        self,
        paginas
    ):

        faces = []

        indice = 0

        while indice < len(paginas):

            pagina_1 = (
                paginas[indice]
                if indice < len(paginas)
                else None
            )

            pagina_2 = (
                paginas[indice + 1]
                if indice + 1 < len(paginas)
                else None
            )

            pagina_3 = (
                paginas[indice + 2]
                if indice + 2 < len(paginas)
                else None
            )

            pagina_4 = (
                paginas[indice + 3]
                if indice + 3 < len(paginas)
                else None
            )

            # ----------------------------------------------
            # FRENTE
            #
            # Página lógica 1 | Página lógica 3
            #
            # Exemplo:
            #
            # 1 | 3
            # 5 | 7
            # 9 | 11
            # ----------------------------------------------

            if (
                pagina_1 is not None
                or pagina_3 is not None
            ):

                faces.append({
                    "esquerda": pagina_1,
                    "direita": pagina_3
                })

            # ----------------------------------------------
            # VERSO
            #
            # Página lógica 4 | Página lógica 2
            #
            # O verso precisa ser espelhado horizontalmente
            # para que, depois da virada manual e do corte:
            #
            # 1 fique atrás de 2
            # 3 fique atrás de 4
            #
            # Exemplo:
            #
            # 4  | 2
            # 8  | 6
            # 12 | 10
            # ----------------------------------------------

            if (
                pagina_2 is not None
                or pagina_4 is not None
            ):

                faces.append({
                    "esquerda": pagina_4,
                    "direita": pagina_2
                })

            indice += 4

        return faces

    # ==================================================
    # RENDERIZA UMA METADE DA FOLHA
    # ==================================================

    def _renderizar_pagina_logica(
        self,
        planilha,
        escala,
        pagina,
        linha_inicio,
        coluna_inicio
    ):

        coluna_fim = (
            coluna_inicio + 3
        )

        preenchimento_titulo = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        preenchimento_periodo = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7"
        )

        preenchimento_cabecalho = PatternFill(
            fill_type="solid",
            fgColor="5B9BD5"
        )

        preenchimento_azul_claro = PatternFill(
            fill_type="solid",
            fgColor="DDEBF7"
        )

        preenchimento_branco = PatternFill(
            fill_type="solid",
            fgColor="FFFFFF"
        )

        lado_fino = Side(
            style="thin",
            color="B7B7B7"
        )

        lado_medio = Side(
            style="medium",
            color="404040"
        )

        # ==================================================
        # CABEÇALHO DA METADE
        # ==================================================

        # --------------------------------------------------
        # CONGREGAÇÃO
        # --------------------------------------------------

        linha = linha_inicio

        planilha.merge_cells(
            start_row=linha,
            start_column=coluna_inicio,
            end_row=linha,
            end_column=coluna_fim
        )

        celula = planilha.cell(
            row=linha,
            column=coluna_inicio
        )

        celula.value = (
            "Congregação Cristã no Brasil"
        )

        celula.font = Font(
            bold=True,
            size=12,
            color="FFFFFF"
        )

        celula.fill = (
            preenchimento_titulo
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha,
            linha_fim=linha,
            coluna_inicio=coluna_inicio,
            coluna_fim=coluna_fim,
            lado=lado_medio
        )

        planilha.row_dimensions[
            linha
        ].height = 18

        # --------------------------------------------------
        # NOME DA ESCALA
        # --------------------------------------------------

        linha += 1

        planilha.merge_cells(
            start_row=linha,
            start_column=coluna_inicio,
            end_row=linha,
            end_column=coluna_fim
        )

        celula = planilha.cell(
            row=linha,
            column=coluna_inicio
        )

        celula.value = escala.nome

        celula.font = Font(
            bold=True,
            size=11,
            color="FFFFFF"
        )

        celula.fill = (
            preenchimento_titulo
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha,
            linha_fim=linha,
            coluna_inicio=coluna_inicio,
            coluna_fim=coluna_fim,
            lado=lado_medio
        )

        planilha.row_dimensions[
            linha
        ].height = 17

        # --------------------------------------------------
        # PERÍODO
        # --------------------------------------------------

        linha += 1

        planilha.merge_cells(
            start_row=linha,
            start_column=coluna_inicio,
            end_row=linha,
            end_column=coluna_fim
        )

        celula = planilha.cell(
            row=linha,
            column=coluna_inicio
        )

        celula.value = (
            f"PERÍODO: "
            f"{self._obter_periodo(escala)}"
        )

        celula.font = Font(
            bold=True,
            size=10
        )

        celula.fill = (
            preenchimento_periodo
        )

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        self._aplicar_borda_intervalo(
            planilha,
            linha_inicio=linha,
            linha_fim=linha,
            coluna_inicio=coluna_inicio,
            coluna_fim=coluna_fim,
            lado=lado_medio
        )

        planilha.row_dimensions[
            linha
        ].height = 16

        # --------------------------------------------------
        # LINHA EM BRANCO
        # --------------------------------------------------

        linha += 1

        planilha.row_dimensions[
            linha
        ].height = 5

        linha += 1

        # ==================================================
        # EVENTOS
        # ==================================================

        if pagina["eventos"]:

            # ----------------------------------------------
            # CABEÇALHO
            # ----------------------------------------------

            cabecalhos = [
                "DATA",
                "EVENTO",
                "ATENDENTE",
                "FUNÇÃO"
            ]

            for deslocamento, titulo in enumerate(
                cabecalhos
            ):

                celula = planilha.cell(
                    row=linha,
                    column=(
                        coluna_inicio
                        + deslocamento
                    )
                )

                celula.value = titulo

                celula.font = Font(
                    bold=True,
                    size=10,
                    color="FFFFFF"
                )

                celula.fill = (
                    preenchimento_cabecalho
                )

                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                celula.border = Border(
                    top=lado_medio,
                    bottom=lado_medio,
                    left=lado_fino,
                    right=lado_fino
                )

            planilha.row_dimensions[
                linha
            ].height = 17

            linha += 1

            # ----------------------------------------------
            # EVENTOS
            # ----------------------------------------------

            for indice_evento, ocorrencia in enumerate(
                pagina["eventos"]
            ):

                linha_inicio_evento = (
                    linha
                )

                atribuicoes = (
                    ocorrencia.atribuicoes
                )

                if atribuicoes:

                    linhas_evento = (
                        len(atribuicoes)
                    )

                else:

                    linhas_evento = 1

                linha_fim_evento = (
                    linha_inicio_evento
                    + linhas_evento
                    - 1
                )

                if indice_evento % 2 == 0:

                    preenchimento = (
                        preenchimento_azul_claro
                    )

                else:

                    preenchimento = (
                        preenchimento_branco
                    )

                if atribuicoes:

                    for atribuicao in atribuicoes:

                        planilha.cell(
                            row=linha,
                            column=coluna_inicio
                        ).value = ocorrencia.data

                        planilha.cell(
                            row=linha,
                            column=coluna_inicio + 1
                        ).value = (
                            ocorrencia.evento.nome
                        )

                        planilha.cell(
                            row=linha,
                            column=coluna_inicio + 2
                        ).value = (
                            atribuicao["pessoa"]
                        )

                        planilha.cell(
                            row=linha,
                            column=coluna_inicio + 3
                        ).value = (
                            atribuicao["funcao"]
                        )

                        linha += 1

                else:

                    planilha.cell(
                        row=linha,
                        column=coluna_inicio
                    ).value = ocorrencia.data

                    planilha.cell(
                        row=linha,
                        column=coluna_inicio + 1
                    ).value = (
                        ocorrencia.evento.nome
                    )

                    linha += 1

                # ------------------------------------------
                # FORMATAÇÃO DAS LINHAS
                # ------------------------------------------

                for linha_evento in range(
                    linha_inicio_evento,
                    linha_fim_evento + 1
                ):

                    for coluna in range(
                        coluna_inicio,
                        coluna_fim + 1
                    ):

                        celula = planilha.cell(
                            row=linha_evento,
                            column=coluna
                        )

                        celula.fill = (
                            preenchimento
                        )

                        celula.font = Font(
                            size=10
                        )

                        celula.alignment = Alignment(
                            horizontal="center",
                            vertical="center"
                        )

                        celula.border = Border(
                            top=lado_fino,
                            bottom=lado_fino,
                            left=lado_fino,
                            right=lado_fino
                        )

                    planilha.cell(
                        row=linha_evento,
                        column=coluna_inicio + 2
                    ).alignment = Alignment(
                        horizontal="left",
                        vertical="center"
                    )

                    planilha.row_dimensions[
                        linha_evento
                    ].height = 17

                # ------------------------------------------
                # DATA
                # ------------------------------------------

                planilha.cell(
                    row=linha_inicio_evento,
                    column=coluna_inicio
                ).number_format = (
                    "dd/mm/yyyy"
                )

                # ------------------------------------------
                # MESCLAR DATA E EVENTO
                # ------------------------------------------

                if (
                    linha_fim_evento
                    > linha_inicio_evento
                ):

                    planilha.merge_cells(
                        start_row=linha_inicio_evento,
                        start_column=coluna_inicio,
                        end_row=linha_fim_evento,
                        end_column=coluna_inicio
                    )

                    planilha.merge_cells(
                        start_row=linha_inicio_evento,
                        start_column=coluna_inicio + 1,
                        end_row=linha_fim_evento,
                        end_column=coluna_inicio + 1
                    )

                planilha.cell(
                    row=linha_inicio_evento,
                    column=coluna_inicio
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                planilha.cell(
                    row=linha_inicio_evento,
                    column=coluna_inicio + 1
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                self._aplicar_borda_externa(
                    planilha,
                    linha_inicio=linha_inicio_evento,
                    linha_fim=linha_fim_evento,
                    coluna_inicio=coluna_inicio,
                    coluna_fim=coluna_fim,
                    lado=lado_medio
                )

        # ==================================================
        # CONTATOS
        # ==================================================

        if pagina["contatos"]:

            if pagina["eventos"]:

                planilha.row_dimensions[
                    linha
                ].height = 5

                linha += 1

            # ----------------------------------------------
            # TÍTULO
            # ----------------------------------------------

            planilha.merge_cells(
                start_row=linha,
                start_column=coluna_inicio,
                end_row=linha,
                end_column=coluna_fim
            )

            celula = planilha.cell(
                row=linha,
                column=coluna_inicio
            )

            celula.value = (
                "CONTATOS DOS ATENDENTES"
            )

            celula.font = Font(
                bold=True,
                size=8,
                color="FFFFFF"
            )

            celula.fill = (
                preenchimento_titulo
            )

            celula.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

            self._aplicar_borda_intervalo(
                planilha,
                linha_inicio=linha,
                linha_fim=linha,
                coluna_inicio=coluna_inicio,
                coluna_fim=coluna_fim,
                lado=lado_medio
            )

            planilha.row_dimensions[
                linha
            ].height = 17

            linha += 1

            # ----------------------------------------------
            # CABEÇALHO
            # ----------------------------------------------

            planilha.merge_cells(
                start_row=linha,
                start_column=coluna_inicio,
                end_row=linha,
                end_column=coluna_inicio + 1
            )

            planilha.merge_cells(
                start_row=linha,
                start_column=coluna_inicio + 2,
                end_row=linha,
                end_column=coluna_fim
            )

            celula_nome = planilha.cell(
                row=linha,
                column=coluna_inicio
            )

            celula_nome.value = "NOME"

            celula_telefone = planilha.cell(
                row=linha,
                column=coluna_inicio + 2
            )

            celula_telefone.value = (
                "TELEFONE"
            )

            for celula in (
                celula_nome,
                celula_telefone
            ):

                celula.font = Font(
                    bold=True,
                    size=7,
                    color="FFFFFF"
                )

                celula.fill = (
                    preenchimento_cabecalho
                )

                celula.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            self._aplicar_borda_intervalo(
                planilha,
                linha_inicio=linha,
                linha_fim=linha,
                coluna_inicio=coluna_inicio,
                coluna_fim=coluna_inicio + 1,
                lado=lado_medio
            )

            self._aplicar_borda_intervalo(
                planilha,
                linha_inicio=linha,
                linha_fim=linha,
                coluna_inicio=coluna_inicio + 2,
                coluna_fim=coluna_fim,
                lado=lado_medio
            )

            planilha.row_dimensions[
                linha
            ].height = 17

            linha += 1

            # ----------------------------------------------
            # PESSOAS
            # ----------------------------------------------

            linha_inicio_contatos = (
                linha
            )

            for indice, pessoa in enumerate(
                pagina["contatos"]
            ):

                planilha.merge_cells(
                    start_row=linha,
                    start_column=coluna_inicio,
                    end_row=linha,
                    end_column=coluna_inicio + 1
                )

                planilha.merge_cells(
                    start_row=linha,
                    start_column=coluna_inicio + 2,
                    end_row=linha,
                    end_column=coluna_fim
                )

                celula_nome = planilha.cell(
                    row=linha,
                    column=coluna_inicio
                )

                celula_telefone = planilha.cell(
                    row=linha,
                    column=coluna_inicio + 2
                )

                celula_nome.value = pessoa.nome
                celula_telefone.value = (
                    pessoa.telefone
                )

                if indice % 2 == 0:

                    preenchimento = (
                        preenchimento_azul_claro
                    )

                else:

                    preenchimento = (
                        preenchimento_branco
                    )

                celula_nome.fill = (
                    preenchimento
                )

                celula_telefone.fill = (
                    preenchimento
                )

                celula_nome.font = Font(
                    size=7
                )

                celula_telefone.font = Font(
                    size=7
                )

                celula_nome.alignment = Alignment(
                    horizontal="left",
                    vertical="center"
                )

                celula_telefone.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                self._aplicar_borda_intervalo(
                    planilha,
                    linha_inicio=linha,
                    linha_fim=linha,
                    coluna_inicio=coluna_inicio,
                    coluna_fim=coluna_inicio + 1,
                    lado=lado_fino
                )

                self._aplicar_borda_intervalo(
                    planilha,
                    linha_inicio=linha,
                    linha_fim=linha,
                    coluna_inicio=coluna_inicio + 2,
                    coluna_fim=coluna_fim,
                    lado=lado_fino
                )

                planilha.row_dimensions[
                    linha
                ].height = 17

                linha += 1

            linha_fim_contatos = (
                linha - 1
            )

            self._aplicar_borda_externa(
                planilha,
                linha_inicio=linha_inicio_contatos - 1,
                linha_fim=linha_fim_contatos,
                coluna_inicio=coluna_inicio,
                coluna_fim=coluna_fim,
                lado=lado_medio
            )

    # ==================================================
    # ALTURA DE UMA FACE FÍSICA
    # ==================================================

    def _altura_face_impressao(
        self
    ):

        # 3 linhas de título
        # 1 linha de separação
        # N linhas úteis de conteúdo
        # 1 linha final de segurança

        return (
            4
            + self._linhas_corpo_impressao
            + 1
        )

    # ==================================================
    # COLUNAS DA IMPRESSÃO
    # ==================================================

    def _ajustar_colunas_impressao(
        self,
        planilha
    ):

        # --------------------------------------------------
        # METADE ESQUERDA
        # --------------------------------------------------

        planilha.column_dimensions[
            "A"
        ].width = 11

        planilha.column_dimensions[
            "B"
        ].width = 10

        planilha.column_dimensions[
            "C"
        ].width = 17

        planilha.column_dimensions[
            "D"
        ].width = 10

        # --------------------------------------------------
        # ESPAÇO PARA CORTE
        # --------------------------------------------------

        planilha.column_dimensions[
            "E"
        ].width = 4

        # --------------------------------------------------
        # METADE DIREITA
        # --------------------------------------------------

        planilha.column_dimensions[
            "F"
        ].width = 11

        planilha.column_dimensions[
            "G"
        ].width = 10

        planilha.column_dimensions[
            "H"
        ].width = 17

        planilha.column_dimensions[
            "I"
        ].width = 10

    # ==================================================
    # CONTATOS ESCALADOS
    # ==================================================

    def _obter_contatos_escalados(
        self,
        escala,
        pessoas
    ):

        nomes_escalados = set()

        for ocorrencia in escala.ocorrencias:

            for atribuicao in (
                ocorrencia.atribuicoes
            ):

                nome = (
                    atribuicao[
                        "pessoa"
                    ].strip()
                )

                if nome:

                    nomes_escalados.add(
                        nome
                    )

        contatos = []

        for pessoa in pessoas:

            if pessoa.nome in nomes_escalados:

                contatos.append(
                    pessoa
                )

        contatos.sort(
            key=lambda pessoa: pessoa.nome
        )

        return contatos

    # ==================================================
    # CONTAGEM DA DISTRIBUIÇÃO
    # ==================================================

    def _contar_distribuicao(
        self,
        ocorrencias
    ):

        contagem = {}

        for ocorrencia in ocorrencias:

            for atribuicao in (
                ocorrencia.atribuicoes
            ):

                pessoa = (
                    atribuicao[
                        "pessoa"
                    ].strip()
                )

                funcao = (
                    atribuicao[
                        "funcao"
                    ].strip()
                )

                if not pessoa:
                    continue

                if pessoa not in contagem:
                    contagem[pessoa] = {}

                if funcao not in contagem[pessoa]:
                    contagem[pessoa][funcao] = 0

                contagem[pessoa][funcao] += 1

        return contagem

    def _obter_funcoes_distribuicao(
        self,
        escala
    ):

        funcoes = []

        for ocorrencia in escala.ocorrencias:

            for atribuicao in (
                ocorrencia.atribuicoes
            ):

                funcao = (
                    atribuicao[
                        "funcao"
                    ].strip()
                )

                if (
                    funcao
                    and funcao not in funcoes
                ):

                    funcoes.append(
                        funcao
                    )

        return funcoes

    def _obter_pessoas_distribuicao(
        self,
        escala
    ):

        return self._obter_pessoas_das_ocorrencias(
            escala.ocorrencias
        )

    def _obter_pessoas_das_ocorrencias(
        self,
        ocorrencias
    ):

        pessoas = set()

        for ocorrencia in ocorrencias:

            for atribuicao in (
                ocorrencia.atribuicoes
            ):

                pessoa = (
                    atribuicao[
                        "pessoa"
                    ].strip()
                )

                if pessoa:
                    pessoas.add(
                        pessoa
                    )

        return sorted(
            pessoas
        )

    def _agrupar_ocorrencias_por_mes(
        self,
        escala
    ):

        meses = {}

        for ocorrencia in escala.ocorrencias:

            chave = (
                ocorrencia.data.year,
                ocorrencia.data.month
            )

            if chave not in meses:
                meses[chave] = []

            meses[chave].append(
                ocorrencia
            )

        return meses

    def _nome_mes(
        self,
        mes
    ):

        nomes = {
            1: "JANEIRO",
            2: "FEVEREIRO",
            3: "MARÇO",
            4: "ABRIL",
            5: "MAIO",
            6: "JUNHO",
            7: "JULHO",
            8: "AGOSTO",
            9: "SETEMBRO",
            10: "OUTUBRO",
            11: "NOVEMBRO",
            12: "DEZEMBRO"
        }

        return nomes[mes]

    # ==================================================
    # BORDAS
    # ==================================================

    def _aplicar_borda_intervalo(
        self,
        planilha,
        linha_inicio,
        linha_fim,
        coluna_inicio,
        coluna_fim,
        lado
    ):

        for linha in range(
            linha_inicio,
            linha_fim + 1
        ):

            for coluna in range(
                coluna_inicio,
                coluna_fim + 1
            ):

                celula = planilha.cell(
                    row=linha,
                    column=coluna
                )

                topo = (
                    lado
                    if linha == linha_inicio
                    else celula.border.top
                )

                inferior = (
                    lado
                    if linha == linha_fim
                    else celula.border.bottom
                )

                esquerda = (
                    lado
                    if coluna == coluna_inicio
                    else celula.border.left
                )

                direita = (
                    lado
                    if coluna == coluna_fim
                    else celula.border.right
                )

                celula.border = Border(
                    top=topo,
                    bottom=inferior,
                    left=esquerda,
                    right=direita
                )

    def _aplicar_borda_externa(
        self,
        planilha,
        linha_inicio,
        linha_fim,
        coluna_inicio,
        coluna_fim,
        lado
    ):

        for linha in range(
            linha_inicio,
            linha_fim + 1
        ):

            celula_esquerda = (
                planilha.cell(
                    row=linha,
                    column=coluna_inicio
                )
            )

            celula_esquerda.border = Border(
                top=celula_esquerda.border.top,
                bottom=celula_esquerda.border.bottom,
                left=lado,
                right=celula_esquerda.border.right
            )

            celula_direita = (
                planilha.cell(
                    row=linha,
                    column=coluna_fim
                )
            )

            celula_direita.border = Border(
                top=celula_direita.border.top,
                bottom=celula_direita.border.bottom,
                left=celula_direita.border.left,
                right=lado
            )

        for coluna in range(
            coluna_inicio,
            coluna_fim + 1
        ):

            celula_superior = (
                planilha.cell(
                    row=linha_inicio,
                    column=coluna
                )
            )

            celula_superior.border = Border(
                top=lado,
                bottom=celula_superior.border.bottom,
                left=celula_superior.border.left,
                right=celula_superior.border.right
            )

            celula_inferior = (
                planilha.cell(
                    row=linha_fim,
                    column=coluna
                )
            )

            celula_inferior.border = Border(
                top=celula_inferior.border.top,
                bottom=lado,
                left=celula_inferior.border.left,
                right=celula_inferior.border.right
            )

    # ==================================================
    # LARGURA DISTRIBUIÇÃO
    # ==================================================

    def _ajustar_largura_distribuicao(
        self,
        planilha,
        ultima_coluna
    ):

        planilha.column_dimensions[
            "A"
        ].width = 28

        for coluna in range(
            2,
            ultima_coluna + 1
        ):

            maior_tamanho = 0

            for linha in range(
                5,
                planilha.max_row + 1
            ):

                valor = planilha.cell(
                    row=linha,
                    column=coluna
                ).value

                if valor is None:
                    continue

                tamanho = len(
                    str(valor)
                )

                if tamanho > maior_tamanho:
                    maior_tamanho = tamanho

            largura = maior_tamanho + 2

            if largura < 10:
                largura = 10

            if largura > 18:
                largura = 18

            planilha.column_dimensions[
                get_column_letter(
                    coluna
                )
            ].width = largura

    # ==================================================
    # LARGURA DA ESCALA
    # ==================================================

    def _ajustar_largura_colunas(
        self,
        planilha
    ):

        for coluna in range(
            1,
            5
        ):

            maior_tamanho = 0

            for linha in range(
                5,
                planilha.max_row + 1
            ):

                celula = planilha.cell(
                    row=linha,
                    column=coluna
                )

                valor = celula.value

                if valor is None:
                    continue

                if (
                    coluna == 1
                    and linha >= 6
                ):

                    tamanho = 10

                else:

                    tamanho = len(
                        str(valor)
                    )

                if tamanho > maior_tamanho:
                    maior_tamanho = tamanho

            largura = (
                maior_tamanho + 2
            )

            if largura > 30:
                largura = 30

            coluna_letra = (
                get_column_letter(
                    coluna
                )
            )

            planilha.column_dimensions[
                coluna_letra
            ].width = largura

    # ==================================================
    # GRUPOS DATA + EVENTO
    # ==================================================

    def _formatar_grupos(
        self,
        planilha,
        lado_medio,
        preenchimento_azul_claro,
        preenchimento_branco
    ):

        if planilha.max_row < 6:
            return

        inicio_grupo = 6
        indice_grupo = 0

        data_anterior = (
            planilha.cell(
                row=6,
                column=1
            ).value
        )

        evento_anterior = (
            planilha.cell(
                row=6,
                column=2
            ).value
        )

        for linha in range(
            7,
            planilha.max_row + 2
        ):

            if linha <= planilha.max_row:

                data_atual = (
                    planilha.cell(
                        row=linha,
                        column=1
                    ).value
                )

                evento_atual = (
                    planilha.cell(
                        row=linha,
                        column=2
                    ).value
                )

            else:

                data_atual = None
                evento_atual = None

            mudou_grupo = (
                data_atual != data_anterior
                or evento_atual
                != evento_anterior
            )

            if mudou_grupo:

                fim_grupo = (
                    linha - 1
                )

                if (
                    indice_grupo % 2
                    == 0
                ):

                    preenchimento = (
                        preenchimento_azul_claro
                    )

                else:

                    preenchimento = (
                        preenchimento_branco
                    )

                self._aplicar_preenchimento_grupo(
                    planilha,
                    inicio_grupo,
                    fim_grupo,
                    preenchimento
                )

                self._aplicar_borda_grupo(
                    planilha,
                    inicio_grupo,
                    fim_grupo,
                    lado_medio
                )

                indice_grupo += 1

                inicio_grupo = linha
                data_anterior = data_atual
                evento_anterior = evento_atual

    def _aplicar_preenchimento_grupo(
        self,
        planilha,
        linha_inicio,
        linha_fim,
        preenchimento
    ):

        for linha in range(
            linha_inicio,
            linha_fim + 1
        ):

            for coluna in range(
                1,
                5
            ):

                planilha.cell(
                    row=linha,
                    column=coluna
                ).fill = preenchimento

    def _aplicar_borda_grupo(
        self,
        planilha,
        linha_inicio,
        linha_fim,
        lado_medio
    ):

        for coluna in range(
            1,
            5
        ):

            celula_superior = (
                planilha.cell(
                    row=linha_inicio,
                    column=coluna
                )
            )

            celula_inferior = (
                planilha.cell(
                    row=linha_fim,
                    column=coluna
                )
            )

            celula_superior.border = Border(
                top=lado_medio,
                bottom=celula_superior.border.bottom,
                left=celula_superior.border.left,
                right=celula_superior.border.right
            )

            celula_inferior.border = Border(
                top=celula_inferior.border.top,
                bottom=lado_medio,
                left=celula_inferior.border.left,
                right=celula_inferior.border.right
            )

        for linha in range(
            linha_inicio,
            linha_fim + 1
        ):

            celula_esquerda = (
                planilha.cell(
                    row=linha,
                    column=1
                )
            )

            celula_direita = (
                planilha.cell(
                    row=linha,
                    column=4
                )
            )

            celula_esquerda.border = Border(
                top=celula_esquerda.border.top,
                bottom=celula_esquerda.border.bottom,
                left=lado_medio,
                right=celula_esquerda.border.right
            )

            celula_direita.border = Border(
                top=celula_direita.border.top,
                bottom=celula_direita.border.bottom,
                left=celula_direita.border.left,
                right=lado_medio
            )

    # ==================================================
    # MESCLAR DATA E EVENTO
    # ==================================================

    def _mesclar_grupos(
        self,
        planilha
    ):

        if planilha.max_row < 6:
            return

        inicio_grupo = 6

        data_anterior = (
            planilha.cell(
                row=6,
                column=1
            ).value
        )

        evento_anterior = (
            planilha.cell(
                row=6,
                column=2
            ).value
        )

        for linha in range(
            7,
            planilha.max_row + 2
        ):

            if linha <= planilha.max_row:

                data_atual = (
                    planilha.cell(
                        row=linha,
                        column=1
                    ).value
                )

                evento_atual = (
                    planilha.cell(
                        row=linha,
                        column=2
                    ).value
                )

            else:

                data_atual = None
                evento_atual = None

            mudou_grupo = (
                data_atual != data_anterior
                or evento_atual
                != evento_anterior
            )

            if mudou_grupo:

                fim_grupo = (
                    linha - 1
                )

                if (
                    fim_grupo
                    > inicio_grupo
                ):

                    planilha.merge_cells(
                        start_row=inicio_grupo,
                        start_column=1,
                        end_row=fim_grupo,
                        end_column=1
                    )

                    planilha.merge_cells(
                        start_row=inicio_grupo,
                        start_column=2,
                        end_row=fim_grupo,
                        end_column=2
                    )

                planilha.cell(
                    row=inicio_grupo,
                    column=1
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                planilha.cell(
                    row=inicio_grupo,
                    column=2
                ).alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                inicio_grupo = linha
                data_anterior = data_atual
                evento_anterior = evento_atual

    # ==================================================
    # ÁREA DE IMPRESSÃO DA ABA ESCALA
    # ==================================================

    def _ajustar_area_impressao(
        self,
        planilha
    ):

        planilha.print_area = (
            f"A1:D{planilha.max_row}"
        )

    # ==================================================
    # PERÍODO
    # ==================================================

    def _obter_periodo(
        self,
        escala
    ):

        if not escala.ocorrencias:
            return "Sem ocorrências"

        primeira_data = (
            escala.ocorrencias[0].data
        )

        ultima_data = (
            escala.ocorrencias[-1].data
        )

        return (
            f"{primeira_data.strftime('%d/%m/%Y')} "
            f"a "
            f"{ultima_data.strftime('%d/%m/%Y')}"
        )

    # ==================================================
    # NOME DO ARQUIVO
    # ==================================================

    def _nome_arquivo(
        self,
        nome
    ):

        caracteres_invalidos = (
            '<>:"/\\|?*'
        )

        nome_limpo = ""

        for caractere in nome:

            if caractere in caracteres_invalidos:

                nome_limpo += "_"

            else:

                nome_limpo += caractere

        nome_limpo = (
            nome_limpo.strip()
        )

        if not nome_limpo:
            nome_limpo = "escala"

        return f"{nome_limpo}.xlsx"